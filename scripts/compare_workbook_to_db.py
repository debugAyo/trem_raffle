import sqlite3
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DB_PATH = ROOT / 'data' / 'raffle.db'
WORKBOOK_PATH = ROOT / 'pre-event-registrations (1).xlsx'


def norm(value):
    return ' '.join(str(value or '').strip().split()).lower()

conn = sqlite3.connect(DB_PATH)
cur = conn.cursor()
existing_names = {norm(name) for (name,) in cur.execute('SELECT name FROM participants').fetchall()}
existing_contacts = {norm(contact) for (contact,) in cur.execute('SELECT contact FROM participants WHERE contact IS NOT NULL').fetchall()}

wb = load_workbook(WORKBOOK_PATH, data_only=True)
ws = wb[wb.sheetnames[0]]
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    record = dict(zip(headers, row))
    if record.get('Full Name'):
        rows.append(record)

missing = []
matched = []
for record in rows:
    name = str(record.get('Full Name') or '').strip()
    phone = str(record.get('Phone') or '').strip()
    email = str(record.get('Email') or '').strip()
    key = norm(name)
    if key in existing_names:
        matched.append(name)
        continue
    if phone and norm(phone) in existing_contacts:
        matched.append(name)
        continue
    if email and norm(email) in existing_contacts:
        matched.append(name)
        continue
    missing.append(name)

print(f'Total workbook rows: {len(rows)}')
print(f'Missing from DB: {len(missing)}')
for name in missing[:30]:
    print('-', name)

print('\nMatched names examples:', matched[:12])
conn.close()
