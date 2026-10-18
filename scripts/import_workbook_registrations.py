from __future__ import annotations

import json
import sqlite3
import uuid
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DB_PATH = ROOT / 'data' / 'raffle.db'
WORKBOOK_PATH = ROOT / 'pre-event-registrations (1).xlsx'


def norm(value):
    return ' '.join(str(value or '').strip().split()).lower()


def read_workbook_rows():
    wb = load_workbook(WORKBOOK_PATH, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row))
        if not record.get('Full Name'):
            continue
        rows.append(record)
    return rows


def choose_contact(record):
    phone = (record.get('Phone') or '').strip()
    email = (record.get('Email') or '').strip()
    if phone:
        return phone
    if email:
        return email
    return None


def generate_unique_raffle_number(conn):
    while True:
        token = f"RFF-{uuid.uuid4().hex[:8].upper()}"
        exists = conn.execute(
            'SELECT 1 FROM participants WHERE raffle_number = ?',
            (token,),
        ).fetchone()
        if not exists:
            return token


def main():
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f'Missing workbook: {WORKBOOK_PATH}')

    rows = read_workbook_rows()
    conn = sqlite3.connect(DB_PATH)
    conn.execute('PRAGMA foreign_keys = ON')

    existing_names = {
        norm(name)
        for (name,) in conn.execute('SELECT name FROM participants').fetchall()
    }
    existing_contacts = {
        norm(contact)
        for (contact,) in conn.execute(
            "SELECT contact FROM participants WHERE contact IS NOT NULL"
        ).fetchall()
    }
    existing_emails = {
        norm(email)
        for (email,) in conn.execute(
            "SELECT contact FROM participants WHERE contact LIKE '%@%'"
        ).fetchall()
    }

    seen_in_workbook = set()
    added = []

    for record in rows:
        name = ' '.join(str(record.get('Full Name') or '').strip().split())
        if not name:
            continue

        key = norm(name)
        if key in seen_in_workbook or key in existing_names:
            continue
        seen_in_workbook.add(key)

        phone = str(record.get('Phone') or '').strip()
        email = str(record.get('Email') or '').strip()
        phone_key = norm(phone)
        email_key = norm(email)
        if phone_key and phone_key in existing_contacts:
            continue
        if email_key and email_key in existing_emails:
            continue

        item_id = f'p-{uuid.uuid4().hex[:8]}'
        card_row = conn.execute(
            'SELECT code FROM card_pool WHERE participant_id IS NULL ORDER BY rowid LIMIT 1'
        ).fetchone()
        if card_row is None:
            raise RuntimeError('No available member cards remain in the card pool.')
        card_code = card_row[0]

        raffle_number = generate_unique_raffle_number(conn)
        contact = choose_contact(record)
        registered_at = datetime.utcnow().isoformat(timespec='seconds') + 'Z'

        conn.execute(
            '''
            INSERT INTO participants (
                id, name, contact, card_code, raffle_number,
                attendance_codes, test_attended, raffle_code,
                invited_by_id, registered_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''',
            (
                item_id,
                name,
                contact,
                card_code,
                raffle_number,
                json.dumps({}),
                0,
                None,
                None,
                registered_at,
            ),
        )
        conn.execute(
            'UPDATE card_pool SET participant_id = ?, assigned_at = ? WHERE code = ? AND participant_id IS NULL',
            (item_id, registered_at, card_code),
        )
        added.append({
            'name': name,
            'card_code': card_code,
            'raffle_number': raffle_number,
            'contact': contact,
        })

    conn.commit()
    conn.close()

    print(f'Workbook rows examined: {len(rows)}')
    print(f'New live registrations imported: {len(added)}')
    for item in added[:10]:
        print(f"- {item['name']} | card={item['card_code']} | raffle={item['raffle_number']}")


if __name__ == '__main__':
    main()
